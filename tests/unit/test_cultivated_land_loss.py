from __future__ import annotations

from pathlib import Path

import pytest
import yaml
from ai_gis_qgis.backend.agent_core import AgentCore
from ai_gis_qgis.backend.executor.artifact_verifier import ArtifactVerifier
from ai_gis_qgis.backend.executor.sandbox_policy import SandboxPolicy
from ai_gis_qgis.backend.skills.skill_manager import SkillManager
from ai_gis_qgis.backend.tools.code_execution import find_generated_code_issues
from ai_gis_qgis.backend.tools.cultivated_land_loss import (
    CLASSIFICATION_PATH,
    REPORT_MAPPING_PATH,
    TEMPLATE_PATH,
    _clone_without_subset,
    _extract_analysis_summary,
    _load_report_mapping,
    _load_rule_bundle,
    _normalize_parameters,
    build_cultivated_land_loss_code,
)
from openpyxl import load_workbook


def _arguments() -> dict[str, object]:
    return {
        "previous_survey_layer_id": "previous-id",
        "increment_layer_id": "increment-id",
        "land_management_layer_id": "management-id",
        "permanent_farmland_layer_id": "permanent-id",
        "previous_land_code_field": "DLBM",
        "increment_land_code_field": "DLBM",
        "management_area_name": "xx区",
        "target_crs": "EPSG:4547",
    }


def test_land_classification_is_complete_and_disjoint():
    rules = _load_rule_bundle()

    assert rules["version"] == "2025-v1"
    assert set(rules["categories"]["cultivated"]) == {"0101", "0102", "0103"}
    assert "05H1" in rules["categories"]["construction"]
    assert "0307" in rules["categories"]["forest_garden"]
    assert "1208" in rules["categories"]["unused"]
    assert len(rules["all_codes"]) == len(set(rules["all_codes"]))
    assert set(rules["code_categories"]) == set(rules["all_codes"])


def test_report_mapping_matches_visible_xx_district_row():
    mapping = _load_report_mapping()

    assert mapping["cells"]["management_area_name"] == "A12"
    assert mapping["cells"]["increment_count"] == "B12"
    assert mapping["cells"]["unreasonable_outflow_area"] == "F12"
    assert mapping["cells"]["current_cultivated_area"] == "S12"
    assert mapping["formulas"]["cultivated_change_area"] == "=F12-P12"
    assert mapping["formulas"]["current_cultivated_area"] == "=Q12-F12+P12"
    assert mapping["unused_template_cells"] == ["D12", "E12"]


def test_xlsx_template_is_packaged_and_matches_report_mapping():
    assert TEMPLATE_PATH.is_file()
    assert TEMPLATE_PATH.suffix == ".xlsx"
    assert TEMPLATE_PATH.read_bytes()[:4] == b"PK\x03\x04"

    verified = ArtifactVerifier().verify(
        {"path": str(TEMPLATE_PATH), "name": "template", "type": "table"}
    )

    assert verified["verified"] is True
    assert verified["metadata"]["format"] == "OOXML/XLSX"
    assert verified["metadata"]["validation_level"] == "structural"
    assert verified["metadata"]["worksheet_count"] == 3

    workbook = load_workbook(TEMPLATE_PATH, data_only=False)
    worksheet = workbook.worksheets[0]
    assert worksheet["A12"].value == "xx区"
    assert worksheet["R12"].value == "=F12-P12"
    assert worksheet["S12"].value == "=Q12-F12+P12"
    assert "A25:S25" in {str(value) for value in worksheet.merged_cells.ranges}


def test_openpyxl_round_trip_preserves_template_layout(tmp_path: Path):
    source = load_workbook(TEMPLATE_PATH, data_only=False)
    source_sheet = source.worksheets[0]
    source_style_id = source_sheet["A12"].style_id
    source_merged = {str(value) for value in source_sheet.merged_cells.ranges}

    source_sheet["A12"] = "测试区"
    source_sheet["B12"] = 123
    source_sheet["R12"] = "=F12-P12"
    source_sheet["S12"] = "=Q12-F12+P12"
    output = tmp_path / "report.xlsx"
    source.save(output)

    result = load_workbook(output, data_only=False)
    result_sheet = result.worksheets[0]
    assert result_sheet["A12"].value == "测试区"
    assert result_sheet["B12"].value == 123
    assert result_sheet["R12"].value == "=F12-P12"
    assert result_sheet["S12"].value == "=Q12-F12+P12"
    assert result_sheet["A12"].style_id == source_style_id
    assert {str(value) for value in result_sheet.merged_cells.ranges} == source_merged


def test_vendored_openpyxl_is_packaged_with_its_dependency():
    root = Path(__file__).resolve().parents[2]
    assert (root / "vendor" / "openpyxl" / "__init__.py").is_file()
    assert (root / "vendor" / "et_xmlfile" / "__init__.py").is_file()


def test_parameter_defaults_preserve_approved_business_formulas():
    parameters = _normalize_parameters(_arguments())

    assert parameters["area_unit"] == "mu"
    assert parameters["analysis_year"] == 2025
    assert parameters["previous_year"] == 2024
    assert parameters["management_area_name"] == "xx区"


@pytest.mark.parametrize(
    "updates,error",
    [
        ({"area_unit": "acre"}, "area_unit"),
        ({"analysis_year": 2024, "previous_year": 2025}, "年份关系"),
        ({"management_area_name": ""}, "management_area_name"),
    ],
)
def test_parameter_validation_rejects_invalid_bindings(updates, error):
    arguments = _arguments()
    arguments.update(updates)

    with pytest.raises(ValueError, match=error):
        _normalize_parameters(arguments)


def test_fixed_script_builds_without_forbidden_runtime_calls(tmp_path: Path):
    arguments = _arguments()
    arguments.update(
        {
            "classification": _load_rule_bundle(),
            "report_mapping": _load_report_mapping(),
            "template_path": str(TEMPLATE_PATH),
        }
    )

    code = build_cultivated_land_loss_code(arguments)
    SandboxPolicy(tmp_path).validate_code(code)
    assert find_generated_code_issues(code) == []
    assert "from PyQt5" not in code
    assert "from PyQt6" not in code

    assert "CULTIVATED_LAND_LOSS_PARAMETERS_JSON" in code
    assert "cultivated_land_loss_analysis.xlsx" in code
    assert "from openpyxl import load_workbook" in code
    assert "QAxContainer" not in code
    assert "Excel.Application" not in code
    assert "=F12-P12" in code
    assert "=Q12-F12+P12" in code
    assert "record_skipped_feature" in code
    assert '"skipped_feature_counts"' in code
    assert "clone_without_subset" in code
    assert "增量包全部有效要素的地类编码均为耕地" in code


def test_active_layer_subset_is_cleared_only_on_analysis_clone():
    class FakeLayer:
        def __init__(self, subset: str):
            self.subset = subset

        def subsetString(self):
            return self.subset

        def clone(self):
            return FakeLayer(self.subset)

        def isValid(self):
            return True

        def setSubsetString(self, value):
            self.subset = value
            return True

    visible = FakeLayer("DLBM IN ('0101','0102','0103')")

    analysis, ignored_subset = _clone_without_subset(visible, "年度增量包")

    assert analysis is not visible
    assert analysis.subsetString() == ""
    assert visible.subsetString() == "DLBM IN ('0101','0102','0103')"
    assert ignored_subset == "DLBM IN ('0101','0102','0103')"


def test_analysis_summary_exposes_non_blocking_data_warnings():
    stdout = (
        "diagnostic\n"
        '{"outputs": ["report.xlsx"], "data_quality": '
        '{"has_warnings": true, "skipped_feature_total": 1}}\n'
    )

    summary = _extract_analysis_summary(stdout)

    assert summary["data_quality"]["has_warnings"] is True
    assert summary["data_quality"]["skipped_feature_total"] == 1


def test_confirmed_analysis_success_is_formatted_as_user_facing_summary():
    result = {
        "analysis_summary": {
            "management_area_name": "汉阳区",
            "analysis_year": 2025,
            "increment_count": 66,
            "unit_label": "亩",
            "metrics": {
                "increment_area": 250.79,
                "unreasonable_outflow_area": 37.58,
                "unreasonable_outflow_permanent_area": 7.97,
                "non_agricultural_area": 17.62,
                "non_agricultural_permanent_area": 1.79,
                "non_grain_area": 19.96,
                "non_grain_permanent_area": 6.18,
                "forest_garden_area": 16.28,
                "forest_garden_permanent_area": 3.28,
                "other_agricultural_area": 3.68,
                "other_agricultural_permanent_area": 2.9,
                "added_cultivated_area": 224.11,
                "previous_cultivated_area": 4198.7,
                "cultivated_change_area": -186.53,
                "current_cultivated_area": 4385.23,
            },
        },
        "data_quality": {
            "warning_count": 1,
            "messages": ["年度增量包图层的活动筛选已在分析副本中忽略"],
        },
        "outputs": [
            {"name": "国土变更调查统计表", "path": "report.xlsx"},
        ],
    }

    core = object.__new__(AgentCore)
    content = core._format_tool_success(
        "execute_cultivated_land_loss_analysis", result
    )

    assert "2025年度国土变更调查耕地流失分析已完成" in content
    assert "耕地不合理流出：37.58 亩" in content
    assert "非农化：17.62 亩" in content
    assert "非粮化：19.96 亩" in content
    assert "数据质量警告：1 条" in content
    assert "report.xlsx" in content
    assert "已确认并执行工具" not in content


def test_skill_routes_only_to_read_only_inspection_and_fixed_execution():
    skills_dir = Path(__file__).resolve().parents[2] / "skills"
    skill = SkillManager(skills_dir).get("analyze-cultivated-land-loss")

    assert skill is not None
    assert skill.platforms == ["linux", "Windows", "macos"]
    assert skill.tools == [
        "list_layers",
        "inspect_layers",
        "inspect_cultivated_land_loss_inputs",
        "get_task_context",
        "execute_cultivated_land_loss_analysis",
    ]
    assert "execute_gis_code" not in skill.tools
    assert "record_pipeline_stage" not in skill.tools


def test_tool_outputs_map_to_skill_completion_artifacts():
    result = {
        "outputs": [
            {"path": "report.xlsx", "name": "report", "verified": True},
            {"path": "details.gpkg", "name": "details", "verified": True},
            {"path": "metrics.csv", "name": "metrics", "verified": True},
            {"path": "quality.json", "name": "quality", "verified": True},
        ]
    }

    artifacts = AgentCore._artifacts_from_tool_result(
        "execute_cultivated_land_loss_analysis", result
    )

    assert [artifact["artifact_type"] for artifact in artifacts] == [
        "report_table",
        "detail_layer",
        "metrics_table",
        "quality_report",
    ]
    assert all(artifact["verified"] for artifact in artifacts)


def test_reference_yaml_files_are_parseable():
    for path in (CLASSIFICATION_PATH, REPORT_MAPPING_PATH):
        payload = yaml.safe_load(path.read_text(encoding="utf-8"))
        assert isinstance(payload, dict)
